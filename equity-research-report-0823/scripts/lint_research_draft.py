#!/usr/bin/env python3
"""Lint a Markdown, text, or DOCX equity-research draft for structural and prose issues."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path


PLACEHOLDER_PATTERNS = (
    re.compile(r"\[\[[^\]]+\]\]"),
    re.compile(r"\b(?:TBD|TODO|FIXME)\b", re.IGNORECASE),
    re.compile(r"(?:待补充|待核实|待确认|此处插入|占位符)"),
)
CAPTION_RE = re.compile(r"^(?:图表|图|表)\s*(?:[FTI]\d{2}|\d+)\s*[：:]", re.MULTILINE)
SOURCE_RE = re.compile(r"(?:资料来源|数据来源|来源)\s*[：:]")
MODAL_RE = re.compile(r"我们认为|我们判断|我们预计|有望|或将|可能")
JUDGMENT_RE = re.compile(r"我们认为|我们判断|我们预计")
ABSOLUTE_RE = re.compile(r"必然|必定|一定会|毫无疑问|确定无疑")
VAGUE_RE = re.compile(r"持续赋能|全面赋能|前景广阔|空间广阔|值得期待|确定性极强|打造闭环")
AUDIT_TAG_RE = re.compile(
    r"【(?:披露事实|已披露事实|样本观察|自行测算|研究测算|研究假设|研究判断|预测|风险情景)】"
)
GENERIC_HEADING_RE = re.compile(
    r"^\s{0,3}#{2,4}\s+(?:\d+(?:\.\d+)*[.、]?\s*)?"
    r"(公司介绍|公司概况|行业情况|行业概况|经营情况|竞争格局|核心优势|核心竞争力|"
    r"产品端|渠道端|营销端|未来空间)\s*$",
    re.MULTILINE,
)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("draft", type=Path)
    parser.add_argument("--json", dest="json_path", type=Path, help="write machine-readable result")
    parser.add_argument("--strict", action="store_true", help="treat warnings as failure")
    parser.add_argument(
        "--fragment",
        action="store_true",
        help="lint a requested section/excerpt without full-report completeness checks",
    )
    return parser.parse_args(argv)


def read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - environment guidance
        raise SystemExit("python-docx is required to lint DOCX files") from exc
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row._tr.tc_lst:
                blocks.append("".join(cell.itertext()))
    for section in document.sections:
        blocks.extend(paragraph.text for paragraph in section.header.paragraphs)
        blocks.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n\n".join(blocks)


def read_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        return read_docx(path)
    if path.suffix.lower() not in {".md", ".markdown", ".txt"}:
        raise SystemExit("supported formats: .md, .markdown, .txt, .docx")
    return path.read_text(encoding="utf-8")


def issue(level: str, code: str, message: str, evidence: str | None = None) -> dict:
    result = {"level": level, "code": code, "message": message}
    if evidence:
        result["evidence"] = evidence
    return result


def find_duplicate_paragraphs(text: str) -> list[str]:
    paragraphs = [
        re.sub(r"\s+", " ", paragraph).strip()
        for paragraph in re.split(r"\n\s*\n", text)
    ]
    eligible = [paragraph for paragraph in paragraphs if len(paragraph) >= 100]
    counts = Counter(eligible)
    return [paragraph for paragraph, count in counts.items() if count > 1]


def caption_source_warnings(text: str) -> list[str]:
    lines = text.splitlines()
    missing = []
    for index, line in enumerate(lines):
        if not CAPTION_RE.search(line.strip()):
            continue
        nearby = "\n".join(lines[index + 1:index + 7])
        if not SOURCE_RE.search(nearby):
            missing.append(line.strip()[:120])
    return missing


def prose_paragraphs(text: str) -> list[str]:
    """Return likely narrative paragraphs while excluding headings, tables, and source notes."""
    paragraphs = []
    for block in re.split(r"\n\s*\n", text):
        compact = re.sub(r"\s+", " ", block).strip()
        if len(compact) < 30:
            continue
        if re.match(
            r"^(?:#{1,6}\s|```|\|\s*|>\s*|[-*+]\s+|\d+[.)、]\s+|"
            r"图表?\s*\d+\s*[：:]|资料来源\s*[：:]|数据来源\s*[：:]|来源\s*[：:]|"
            r"报告类型\s*[：:]|数据截止日\s*[：:]|估值基准日\s*[：:]|报告币种)",
            compact,
        ):
            continue
        paragraphs.append(compact)
    return paragraphs


def sentence_units(paragraph: str) -> list[str]:
    return [
        unit.strip()
        for unit in re.split(r"(?<=[。！？；])", paragraph)
        if unit.strip()
    ]


def style_warnings(text: str) -> list[dict]:
    """Flag prose patterns that weaken the sampled sell-side research style."""
    warnings = []
    paragraphs = prose_paragraphs(text)
    sentences = [unit for paragraph in paragraphs for unit in sentence_units(paragraph)]

    long_paragraphs = [paragraph for paragraph in paragraphs if len(paragraph) > 420]
    if long_paragraphs:
        warnings.append(
            issue(
                "warning",
                "LONG_PARAGRAPH",
                f"发现 {len(long_paragraphs)} 个超过 420 字的正文段落；检查是否包含多个段首标签、独立结论或认识层级换挡。",
                long_paragraphs[0][:180],
            )
        )

    long_sentences = [sentence for sentence in sentences if len(sentence) > 120]
    if long_sentences:
        warnings.append(
            issue(
                "warning",
                "LONG_SENTENCE",
                f"发现 {len(long_sentences)} 个超过 120 字的句段；优先在因果或并列层级处断开。",
                long_sentences[0][:180],
            )
        )

    modal_stacks = [
        (paragraph, len(MODAL_RE.findall(paragraph)))
        for paragraph in paragraphs
        if len(MODAL_RE.findall(paragraph)) >= 4
    ]
    if modal_stacks:
        paragraph, count = modal_stacks[0]
        warnings.append(
            issue(
                "warning",
                "MODAL_STACKING",
                f"发现 {len(modal_stacks)} 个判断/预测标记过密的段落；首个段落含 {count} 处。",
                paragraph[:180],
            )
        )

    repeated_modals = []
    for sentence in sentences:
        clauses = re.split(r"[；;]|[，,](?:但|而|另一方面|同时)", sentence)
        if any(len(re.findall(r"有望|或将|可能", clause)) >= 2 for clause in clauses):
            repeated_modals.append(sentence)
    if repeated_modals:
        warnings.append(
            issue(
                "warning",
                "REPEATED_PROSPECTIVE_MODAL",
                "同一句叠加多个‘有望/或将/可能’；保留一个标记并补充成立条件。",
                repeated_modals[0][:180],
            )
        )

    absolute_matches = [sentence for sentence in sentences if ABSOLUTE_RE.search(sentence)]
    if absolute_matches:
        warnings.append(
            issue(
                "warning",
                "OVERCONFIDENT_WORDING",
                "发现无条件确定性表述；改为可核验事实或带条件的预测。",
                absolute_matches[0][:180],
            )
        )

    vague_matches = [sentence for sentence in sentences if VAGUE_RE.search(sentence)]
    if vague_matches:
        warnings.append(
            issue(
                "warning",
                "VAGUE_BOILERPLATE",
                "发现空泛研报套话；用具体变量、传导科目和验证指标替换。",
                vague_matches[0][:180],
            )
        )

    audit_tags = AUDIT_TAG_RE.findall(text)
    if len(audit_tags) >= 2:
        warnings.append(
            issue(
                "warning",
                "AUDIT_LABEL_IN_PROSE",
                "成稿反复使用证据审计标签；用‘根据公司披露/样本显示/在假设下’自然融入句子。",
                "、".join(dict.fromkeys(audit_tags)),
            )
        )

    prose_text = "".join(paragraphs)
    judgment_count = len(JUDGMENT_RE.findall(prose_text))
    density = judgment_count * 1000 / max(len(prose_text), 1)
    if judgment_count >= 5 and density > 3:
        warnings.append(
            issue(
                "warning",
                "JUDGMENT_OVERUSE",
                f"‘我们认为/判断/预计’使用偏密（{density:.1f} 次/千字）；让事实和标题承担更多方向表达。",
            )
        )

    repetitive_openers = []
    openers = ("我们认为", "我们判断", "我们预计", "同时", "此外", "其中")
    for index in range(len(paragraphs) - 2):
        window = paragraphs[index:index + 3]
        for opener in openers:
            if all(paragraph.startswith(opener) for paragraph in window):
                repetitive_openers.append((opener, window[0]))
                break
    if repetitive_openers:
        opener, evidence = repetitive_openers[0]
        warnings.append(
            issue(
                "warning",
                "REPETITIVE_OPENERS",
                f"至少三个相邻段落均以‘{opener}’起笔；改用对象、时间、数据或主题标签切入。",
                evidence[:180],
            )
        )

    generic_headings = [match.group(0).strip() for match in GENERIC_HEADING_RE.finditer(text)]
    if generic_headings:
        warnings.append(
            issue(
                "warning",
                "GENERIC_HEADING",
                f"发现 {len(generic_headings)} 个只有主题、没有结论的标题；补充阶段或方向性判断。",
                generic_headings[0],
            )
        )

    return warnings


def lint(text: str, fragment: bool = False) -> list[dict]:
    issues = []
    normalized = re.sub(r"\s+", " ", text)

    required = {
        "MISSING_AS_OF": (r"(?:数据|资料)(?:截止|截至)(?:日|日期)?|估值基准日|\bas of\b", "缺少数据截止日或估值基准日。"),
        "MISSING_CONCLUSION": (r"核心观点|投资要点|研究结论|核心结论|核心投资逻辑|投资建议|投资策略", "缺少可识别的研究结论/核心观点部分。"),
        "MISSING_RISKS": (r"风险因素|风险提示|主要风险|证伪条件", "缺少风险或证伪条件部分。"),
        "MISSING_SOURCES": (r"资料来源|数据来源|参考资料|来源[：:]|\bSources?\b", "缺少来源标注。"),
    }
    if not fragment:
        for code, (pattern, message) in required.items():
            if not re.search(pattern, normalized, re.IGNORECASE):
                issues.append(issue("blocker", code, message))

    placeholders = []
    for pattern in PLACEHOLDER_PATTERNS:
        placeholders.extend(match.group(0) for match in pattern.finditer(text))
    if placeholders:
        sample = ", ".join(dict.fromkeys(placeholders[:8]))
        issues.append(issue("blocker", "PLACEHOLDER", "草稿仍含未完成占位符。", sample))

    for caption in caption_source_warnings(text):
        issues.append(
            issue("warning", "FIGURE_WITHOUT_SOURCE", "图表附近未找到独立来源标注。", caption)
        )

    duplicates = find_duplicate_paragraphs(text)
    if duplicates:
        issues.append(
            issue(
                "warning",
                "DUPLICATE_PARAGRAPH",
                f"发现 {len(duplicates)} 个完全重复的长段落。",
                duplicates[0][:160],
            )
        )

    issues.extend(style_warnings(text))

    if not fragment:
        if re.search(r"盈利预测|业绩预测|\d{4}E", normalized, re.IGNORECASE) and not re.search(
            r"关键假设|预测假设|假设如下|驱动", normalized
        ):
            issues.append(issue("warning", "FORECAST_WITHOUT_ASSUMPTIONS", "出现预测结果，但未找到关键假设说明。"))

        if re.search(r"估值|目标价|评级", normalized) and not re.search(
            r"可比|PE|P/E|EV/EBITDA|DCF|SOTP|分部估值|现金流折现", normalized, re.IGNORECASE
        ):
            issues.append(issue("warning", "VALUATION_WITHOUT_METHOD", "出现估值/评级，但未找到估值方法。"))

        if re.search(r"目标价", normalized) and not re.search(r"价格基准日|估值基准日|收盘价", normalized):
            issues.append(issue("warning", "TARGET_WITHOUT_PRICE_DATE", "出现目标价，但未找到价格基准日。"))

        if len(normalized) < 1200:
            issues.append(issue("warning", "VERY_SHORT", "草稿较短；确认这与所选报告形态一致。"))

    return issues


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    path = args.draft.expanduser().resolve()
    if not path.is_file():
        raise SystemExit(f"draft not found: {path}")
    issues = lint(read_text(path), fragment=args.fragment)
    result = {
        "draft": str(path),
        "blockers": sum(item["level"] == "blocker" for item in issues),
        "warnings": sum(item["level"] == "warning" for item in issues),
        "issues": issues,
    }
    for item in issues:
        evidence = f" — {item['evidence']}" if item.get("evidence") else ""
        print(f"[{item['level'].upper()}] {item['code']}: {item['message']}{evidence}")
    if not issues:
        print("PASS: no structural blockers or warnings found")
    else:
        print(f"SUMMARY: {result['blockers']} blocker(s), {result['warnings']} warning(s)")
    if args.json_path:
        args.json_path.parent.mkdir(parents=True, exist_ok=True)
        args.json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return 1 if result["blockers"] or (args.strict and result["warnings"]) else 0


if __name__ == "__main__":
    raise SystemExit(main())
