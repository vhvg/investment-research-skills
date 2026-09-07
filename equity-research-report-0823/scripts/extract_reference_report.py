#!/usr/bin/env python3
"""Extract reusable structural evidence from one or more research-report DOCX files."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError as exc:  # pragma: no cover - environment guidance
    raise SystemExit("python-docx is required; use the bundled document workspace runtime") from exc


HEADING_RE = re.compile(
    r"^(?:第[一二三四五六七八九十百]+[章节篇]|[一二三四五六七八九十]+[、.]|"
    r"\d+(?:\.\d+){0,3}[、.．\s]|核心观点|投资要点|研究结论|投资建议|"
    r"盈利预测|估值|风险提示|风险因素|目录|摘要|附录|行业观点)"
)
SOURCE_RE = re.compile(r"^(?:资料来源|数据来源|来源)[：:]")
CAPTION_RE = re.compile(r"^(?:图表|图|表)\s*\d+\s*[：:]")
APP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("reports", nargs="+", type=Path, help="DOCX report paths")
    parser.add_argument("--out-dir", required=True, type=Path, help="directory for JSON evidence")
    parser.add_argument(
        "--include-text",
        action="store_true",
        help="include all top-level paragraph text; off by default to keep evidence compact",
    )
    return parser.parse_args(argv)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def cm(value) -> float | None:
    return round(value.cm, 3) if value is not None else None


def outline_level(paragraph):
    properties = paragraph._p.pPr
    if properties is None:
        return None
    element = properties.find(qn("w:outlineLvl"))
    if element is None:
        return None
    value = element.get(qn("w:val"))
    return int(value) if value and value.isdigit() else value


def font_summary(paragraph) -> dict:
    sizes: list[float] = []
    names: set[str] = set()
    colors: set[str] = set()
    bold_chars = 0
    total_chars = 0
    for run in paragraph.runs:
        length = len(run.text.strip())
        total_chars += length
        if run.bold:
            bold_chars += length
        if run.font.size:
            sizes.append(round(run.font.size.pt, 2))
        if run.font.name:
            names.add(run.font.name)
        if run.font.color and run.font.color.rgb:
            colors.add(str(run.font.color.rgb))
    return {
        "max_size_pt": max(sizes) if sizes else None,
        "min_size_pt": min(sizes) if sizes else None,
        "font_names": sorted(names),
        "colors": sorted(colors),
        "bold_ratio": round(bold_chars / total_chars, 3) if total_chars else 0,
    }


def table_summary(table, index: int) -> dict:
    preview = []
    for row in table.rows[:3]:
        values = []
        # Read low-level cells because some broker templates contain malformed
        # vertical merges that python-docx cannot expose through row.cells.
        for cell in row._tr.tc_lst:
            text = re.sub(r"\s+", " ", "".join(cell.itertext())).strip()
            values.append(text[:200])
        preview.append(values)
    return {
        "index": index,
        "rows": len(table.rows),
        "cols": max((len(row._tr.tc_lst) for row in table.rows), default=0),
        "style": table.style.name if table.style else None,
        "preview": preview,
    }


def package_evidence(path: Path) -> dict:
    evidence = {
        "stored_page_count": None,
        "media_parts": 0,
        "header_parts": 0,
        "footer_parts": 0,
        "has_footnotes": False,
        "has_comments": False,
        "custom_xml_parts": 0,
    }
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        evidence.update(
            media_parts=sum(name.startswith("word/media/") for name in names),
            header_parts=sum(
                name.startswith("word/header") and name.endswith(".xml") for name in names
            ),
            footer_parts=sum(
                name.startswith("word/footer") and name.endswith(".xml") for name in names
            ),
            has_footnotes="word/footnotes.xml" in names,
            has_comments="word/comments.xml" in names,
            custom_xml_parts=sum(name.startswith("customXml/") for name in names),
        )
        if "docProps/app.xml" in names:
            root = ElementTree.fromstring(archive.read("docProps/app.xml"))
            pages = root.find(f"{{{APP_NS}}}Pages")
            if pages is not None and pages.text and pages.text.isdigit():
                evidence["stored_page_count"] = int(pages.text)
    return evidence


def analyze(path: Path, include_text: bool) -> dict:
    document = Document(path)
    styles = Counter()
    fonts = Counter()
    sizes = Counter()
    headings = []
    captions = []
    sources = []
    paragraphs = []

    for index, paragraph in enumerate(document.paragraphs):
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        styles[style] += 1
        font = font_summary(paragraph)
        for name in font["font_names"]:
            fonts[name] += len(text)
        if font["max_size_pt"]:
            sizes[str(font["max_size_pt"])] += 1
        level = outline_level(paragraph)
        record = {
            "index": index,
            "style": style,
            "outline_level": level,
            "text": text,
            **font,
        }
        if include_text:
            paragraphs.append(record)
        heading_style = any(token in style.lower() for token in ("heading", "标题", "title"))
        if level is not None or heading_style or HEADING_RE.search(text):
            headings.append(record)
        if CAPTION_RE.search(text):
            captions.append(record)
        if SOURCE_RE.search(text):
            sources.append(record)

    sections = []
    for index, section in enumerate(document.sections):
        sections.append(
            {
                "index": index,
                "page_width_cm": cm(section.page_width),
                "page_height_cm": cm(section.page_height),
                "top_margin_cm": cm(section.top_margin),
                "bottom_margin_cm": cm(section.bottom_margin),
                "left_margin_cm": cm(section.left_margin),
                "right_margin_cm": cm(section.right_margin),
                "header_distance_cm": cm(section.header_distance),
                "footer_distance_cm": cm(section.footer_distance),
                "start_type": str(section.start_type),
                "different_first_page": section.different_first_page_header_footer,
            }
        )

    result = {
        "path": str(path),
        "sha256": sha256(path),
        "file_size": path.stat().st_size,
        "top_level_paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "sections": sections,
        "package": package_evidence(path),
        "style_counts": styles.most_common(30),
        "font_char_counts": fonts.most_common(20),
        "direct_size_counts": sizes.most_common(20),
        "heading_candidates": headings,
        "caption_candidates": captions,
        "source_lines": sources,
        "tables": [table_summary(table, index) for index, table in enumerate(document.tables)],
    }
    if include_text:
        result["paragraphs"] = paragraphs
    return result


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    args.out_dir.mkdir(parents=True, exist_ok=True)
    summaries = []
    for raw_path in args.reports:
        path = raw_path.expanduser().resolve()
        if path.suffix.lower() != ".docx" or not path.is_file():
            raise SystemExit(f"not a readable DOCX: {path}")
        result = analyze(path, args.include_text)
        destination = args.out_dir / f"{path.stem}.json"
        destination.write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        summaries.append(
            {
                key: value
                for key, value in result.items()
                if key not in {"paragraphs", "tables", "heading_candidates", "caption_candidates"}
            }
        )
        print(
            f"{path.name}: {len(result['heading_candidates'])} heading candidates, "
            f"{result['table_count']} tables, {result['package']['media_parts']} media parts"
        )
    (args.out_dir / "summary.json").write_text(
        json.dumps(summaries, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
